"""Word（.docx）来源测试：标题层级→section、列表、表格→Markdown+JSON、身份与消歧、
缓存、脱敏、分块复用。

测试用 python-docx **现场生成** docx（不塞二进制夹具）：夹具二进制无法 review，
且生成方式本身就是"我们支持哪些写法"的可执行文档。
"""
import json
import os
import shutil
import struct
import tempfile
import unittest
import zlib
from pathlib import Path
from unittest import mock

from vllm_kb.config import AppConfig, SourceCfg
from vllm_kb.sources import WordSource


def make_png(w=8, h=8, color=(255, 0, 0)) -> bytes:
    """纯 zlib 造一张 PNG（不依赖 PIL）：夹具要够小、可 review，且内容可区分（测内容寻址）。"""
    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xffffffff)

    raw = b"".join(b"\x00" + bytes(color) * w for _ in range(h))
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def make_docx(path: Path, *, headings=True, table=True, lists=True,
              note="现场：训练任务在第 300 步卡住。") -> None:
    """生成含标题/正文/列表/表格/空段落的测试 docx（覆盖中英样式名）。

    `headings=False` 时仍写一段正文——空文档会被来源跳过，用它构造"改过内容"的版本。
    """
    import docx

    d = docx.Document()
    if headings:
        d.add_heading("HCCL 超时排查案例", level=1)
        d.add_paragraph(note)
        d.add_heading("现象", level=2)
        d.add_paragraph("节点 10.0.0.5 报错，日志在 /home/user/logs/hccl.log。")
    else:
        d.add_paragraph(note)
    if lists:
        d.add_heading("排查步骤", level=2)
        d.add_paragraph("检查网卡状态", style="List Bullet")
        d.add_paragraph("确认 rank 表一致", style="List Number")
    if table:
        d.add_heading("错误码表", level=2)
        t = d.add_table(rows=3, cols=3)
        data = [["错误码", "含义", "处理"],
                ["561000", "aclnn 执行失败", "检查算子支持"],
                ["107020", "dispatch_ffn_combine 失败", "升级驱动"]]
        for r, row in enumerate(data):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    d.add_paragraph("")          # 空段落应跳过
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


def make_docx_with_images(path: Path, *, n=1, alt=None, colors=None, orphan=False,
                          only_image=False, heading=True) -> None:
    """造含**内嵌图片**的 docx（现场生成，不塞二进制夹具）。

    - `n`：正文里的图片张数；`colors` 决定每张图的内容（不同内容 → 不同 sha，测内容寻址）；
    - `alt`：写进 `wp:docPr/@descr`（→ `[图片:alt]` 占位）；
    - `orphan=True`：插一张图后把 `w:drawing` 摘掉——**关系仍在 rels 里**，于是成为孤儿图；
    - `only_image=True`：只放一张图、不放任何文字（回归：以前这种文档会被整篇跳过）。
    """
    import io

    import docx
    from docx.oxml.ns import qn

    d = docx.Document()
    if not only_image and heading:
        d.add_heading("图片案例", level=1)
        d.add_paragraph("下图是拓扑：")
    colors = list(colors or [(255, 0, 0), (0, 128, 255), (0, 200, 0)])
    for i in range(n):
        if not only_image:
            d.add_paragraph(f"图 {i + 1}：")
        run = d.add_paragraph().add_run()
        run.add_picture(io.BytesIO(make_png(color=colors[i % len(colors)])))
        if alt:
            run._r.xpath(".//wp:docPr")[0].set("descr", alt)
    if orphan:
        run = d.add_paragraph().add_run()
        run.add_picture(io.BytesIO(make_png(color=(9, 9, 9))))
        drawing = run._r.xpath(".//w:drawing")[0]
        drawing.getparent().remove(drawing)      # 摘掉引用，关系留在 rels → 孤儿
    if not only_image and heading:
        d.add_paragraph("结尾正文")
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


def _docx_rId(path: Path, index: int = 0) -> str:
    """取第 index 张图片的 rId（造 VML / AlternateContent 用例要用真实关系 id）。"""
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(str(path))
    for par in d.paragraphs:
        blips = par._p.xpath(".//a:blip")
        if blips:
            return blips[index].get(qn("r:embed"))
    raise AssertionError("文档里没有图片")


class _WordCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.import_dir = self.root / "data" / "imports" / "word"
        self.import_dir.mkdir(parents=True)
        self.cfg = SourceCfg(id="cases", type="word", path="data/imports/word", enabled=True)

    def tearDown(self):
        os.environ.pop("VLLM_KB_DATA_ROOT", None)
        self.tmp.cleanup()

    def _src(self, **kw) -> WordSource:
        return WordSource(self.cfg, project_root=self.root, **kw)

    def _write(self, name: str, **kw) -> Path:
        p = self.import_dir / name
        make_docx(p, **kw)
        return p

    def _docs(self, **kw):
        """构建：导入目录还没有 docx 时自动放一篇默认案例（多数用例只关心解析结果）。"""
        if self.import_dir.exists() and not any(self.import_dir.rglob("*.docx")):
            self._write("案例.docx", **kw)
        src = self._src()
        src.pull()
        return src.canonicalize()


class TestWordSource(_WordCase):
    def test_pull_copies_to_assets(self):
        self._write("案例.docx")
        src = self._src()
        self.assertEqual(src.pull(), 1)
        self.assertTrue((self.root / "data" / "assets" / "word" / "案例.docx").exists())
        self.assertEqual(src.pull(), 0)          # 幂等

    def test_headings_become_markdown_levels(self):
        """标题样式 → `#` 层级（这是分块/标签共用的结构来源）。"""
        docs = self._docs()
        self.assertEqual(len(docs), 1)
        body = docs[0].body
        self.assertIn("# HCCL 超时排查案例", body)
        self.assertIn("## 现象", body)
        self.assertIn("## 排查步骤", body)
        # 一级标题作 title（不是文件名）
        self.assertEqual(docs[0].title, "HCCL 超时排查案例")

    def test_lists_get_prefix(self):
        """列表项加 `- `/`1. ` 前缀——否则条目会粘成一段，FTS 命中粒度差。"""
        body = self._docs()[0].body
        self.assertIn("- 检查网卡状态", body)
        self.assertIn("1. 确认 rank 表一致", body)

    def test_table_into_body_and_json(self):
        """表格既拼进正文（可 FTS 检索）又落结构化 JSON（供图/查询）。"""
        docs = self._docs()
        body = docs[0].body
        self.assertIn("| 错误码 | 含义 | 处理 |", body)
        self.assertIn("| 561000 | aclnn 执行失败 | 检查算子支持 |", body)
        self.assertEqual(docs[0].extra["asset"]["tables"], 1)
        rel = docs[0].extra["structure"]["tables"]
        self.assertEqual(len(rel), 1)
        # 落盘文件名以 asset_id 命名（不暴露文件名/路径）
        aid = docs[0].extra["asset"]["asset_id"]
        self.assertEqual(rel[0], f"parsed/word/{aid}.tables.json")
        tpath = self.root / "data" / rel[0]
        self.assertTrue(tpath.exists())
        data = json.loads(tpath.read_text(encoding="utf-8"))
        self.assertEqual(data["tables"][0]["rows"][1][0], "561000")
        self.assertEqual(data["source"], f"word:{aid}")

    def test_identity_and_metadata(self):
        docs = self._docs()
        d = docs[0]
        self.assertEqual(d.source_id, "word:案例")
        self.assertEqual(d.source_type, "doc_word")
        self.assertEqual(d.extra["asset"]["format"], "word")
        self.assertEqual(d.extra["quality"]["parsed_with"], "python-docx")
        self.assertEqual(d.extra["quality"]["text_source"], "text_layer")
        # 与 markdown/excel 统一路径：先入库，审核工作台补标
        self.assertEqual(d.extra["verification"], "unverified")
        self.assertGreater(d.extra["asset"]["paragraphs"], 0)

    def test_entities_reuse_existing_pipeline(self):
        """表格里的错误码复用现有提取线路（canonical → MENTIONS）。"""
        from vllm_kb.graph_rels import extract_doc_relations

        d = self._docs()[0]
        ex = extract_doc_relations(d.source_id, "", 0, d.source_type, d.body)
        self.assertIn("561000", ex.mentions.get("error_code", set()))

    def test_empty_document_skipped(self):
        """无正文（只含空段落）→ 跳过，不产出空文档。"""
        import docx

        d = docx.Document()
        d.add_paragraph("")
        (self.import_dir / "空.docx").parent.mkdir(parents=True, exist_ok=True)
        d.save(str(self.import_dir / "空.docx"))
        self.assertEqual(self._docs(), [])

    def test_corrupt_file_skipped_without_crash(self):
        """非 OOXML（伪装的 .docx）→ 打印原因并跳过，不影响其它文件。"""
        self._write("好的.docx")
        (self.import_dir / "坏的.docx").write_bytes(b"\xd0\xcf\x11\xe0 not a zip")
        docs = self._docs()
        self.assertEqual([d.source_id for d in docs], ["word:好的"])

    def test_import_dir_wins_over_assets(self):
        """优先读导入目录：改源文件后重跑，仍是同一篇（不会因资产层副本变成两篇）。"""
        p = self._write("案例.docx")
        self._docs()
        # 改内容 → 资产层会另存 案例.<sha12>.docx（不可变层），但 canonicalize 读导入目录
        make_docx(p, headings=False, table=False, lists=False)
        docs = self._docs()
        self.assertEqual([d.source_id for d in docs], ["word:案例"])
        self.assertNotIn("## 现象", docs[0].body)

    def test_fallback_to_assets_collapses_versions(self):
        """导入目录缺失 → 回退资产层，并按版本族收敛（只取最新一份，id 取族名）。"""
        self._write("案例.docx")
        self._docs()
        shutil.rmtree(self.import_dir)
        assets = self.root / "data" / "assets" / "word"
        # 伪造一个"更新版本"（_copy_asset 的同名异内容命名）
        newer = assets / "案例.abcdef123456.docx"
        make_docx(newer, headings=True, table=False, lists=False)
        os.utime(assets / "案例.docx", (1000, 1000))
        os.utime(newer, (2000, 2000))
        docs = self._docs()
        self.assertEqual([d.source_id for d in docs], ["word:案例"])   # 收敛为一篇
        self.assertIn("# HCCL 超时排查案例", docs[0].body)             # 取最新版本

    def test_same_stem_different_dir_disambiguated(self):
        """同名不同目录 → 相对路径指纹消歧（导入目录保留目录树的意义所在）。"""
        make_docx(self.import_dir / "a" / "案例.docx", headings=False, table=False, lists=False)
        make_docx(self.import_dir / "b" / "案例.docx", headings=False, table=False, lists=False)
        ids = sorted(d.source_id for d in self._docs())
        self.assertEqual(len(ids), 2)
        self.assertTrue(all(i.startswith("word:案例--") for i in ids), ids)
        self.assertNotEqual(ids[0], ids[1])

    def test_parse_cache_hit(self):
        """解析中间产物按 asset_id 缓存；内容不变时复用（并写 extract.json）。"""
        self._write("案例.docx")
        self._docs()
        aid = None
        for f in (self.root / "data" / "parsed" / "word").glob("*.extract.json"):
            aid = f.stem.split(".")[0]
            data = json.loads(f.read_text(encoding="utf-8"))
            self.assertEqual(data["asset_id"], aid)
        self.assertIsNotNone(aid, "未写 extract.json 缓存")
        # 二次解析：结果一致（缓存命中路径与首次共用 _doc_from_extract）
        first = self._docs()
        second = self._docs()
        self.assertEqual(first[0].body, second[0].body)
        self.assertEqual(first[0].source_id, second[0].source_id)

    def test_raw_text_stored_and_sanitize_log_collected(self):
        """后置脱敏：body 原文入库（原文检索）；会被脱敏的 IP/路径落维护日志。"""
        os.environ["VLLM_KB_DATA_ROOT"] = str(self.root / "data")
        cfg = AppConfig.model_validate({})
        self._write("案例.docx")
        src = self._src(app_cfg=cfg)
        src.pull()
        docs = src.canonicalize()
        body = docs[0].body
        self.assertIn("10.0.0.5", body)                       # 原文入库
        self.assertIn("/home/user/logs/hccl.log", body)
        log = self.root / "data" / "sanitize_log.json"
        self.assertTrue(log.exists())
        data = json.loads(log.read_text(encoding="utf-8"))
        self.assertIn("10.0.0.5", data["ips"])
        self.assertIn("/home/user/logs/hccl.log", data["paths"])

    def test_asset_registry_registration(self):
        """资产注册到 asset_registry（审核台据此反查路径/下载原件）。"""
        from vllm_kb.review import list_assets

        os.environ["VLLM_KB_DATA_ROOT"] = str(self.root / "data")
        cfg = AppConfig.model_validate({})
        self._write("案例.docx")
        src = self._src(app_cfg=cfg)
        src.pull()
        assets = list_assets(self.root / "data" / "review.sqlite3")
        self.assertEqual(len(assets), 1)
        reg = next(iter(assets.values()))
        self.assertEqual(reg["rel_path"], "assets/word/案例.docx")
        self.assertEqual(reg["source_type"], "doc_word")
        self.assertGreater(reg["size"], 0)

    def test_word_missing_dependency_degrades(self):
        """未装 python-docx → 提示安装并返回空（与 openpyxl/pymupdf 同样的降级风格）。"""
        import builtins

        real_import = builtins.__import__

        def fake_import(name, *a, **kw):
            if name == "docx":
                raise ImportError("no module named docx")
            return real_import(name, *a, **kw)

        self._write("案例.docx")
        with mock.patch.object(builtins, "__import__", side_effect=fake_import):
            self.assertEqual(self._src().canonicalize(), [])


class TestWordChunking(_WordCase):
    def test_chunks_carry_section(self):
        """doc_word 复用 markdown 章节切分：chunk 带 section（命中能看到所属章节）。"""
        from vllm_kb.chunking import chunk_doc

        d = self._docs()[0]
        chunks = chunk_doc(d)
        sections = [c.section for c in chunks]
        self.assertIn("现象", sections)
        self.assertIn("排查步骤", sections)
        # section 前缀注入正文（检索命中直接可读）
        self.assertTrue(any("【现象】" in c.text for c in chunks))


class TestWordHelpers(unittest.TestCase):
    """样式识别与文档顺序遍历的单测（不依赖真实 Word 文件）。"""

    def test_heading_level_english_and_chinese(self):
        import docx

        from vllm_kb.sources import _word_heading_level

        d = docx.Document()
        h1 = d.add_heading("一", level=1)
        h3 = d.add_heading("三", level=3)
        plain = d.add_paragraph("正文")
        self.assertEqual(_word_heading_level(h1), 1)
        self.assertEqual(_word_heading_level(h3), 3)
        self.assertEqual(_word_heading_level(plain), 0)
        # 中文界面样式名（"标题 1"）+ 无 style_id 的兜底路径
        with mock.patch.object(type(h1.style), "name", "标题 2"):
            self.assertEqual(_word_heading_level(h1), 2)

    def test_title_style_is_level_1(self):
        import docx

        from vllm_kb.sources import _word_heading_level

        d = docx.Document()
        p = d.add_paragraph("文档标题", style="Title")
        self.assertEqual(_word_heading_level(p), 1)

    def test_list_prefix(self):
        import docx

        from vllm_kb.sources import _word_list_prefix

        d = docx.Document()
        self.assertEqual(_word_list_prefix(d.add_paragraph("正文")), "")
        self.assertEqual(_word_list_prefix(d.add_paragraph("项", style="List Bullet")), "- ")
        self.assertEqual(_word_list_prefix(d.add_paragraph("项", style="List Number")), "1. ")

    def test_iter_docx_blocks_preserves_order(self):
        """段落与表格必须按文档顺序产出（否则表格会全被挪到正文末尾、section 归属错）。"""
        import docx

        from vllm_kb.sources import _iter_docx_blocks

        d = docx.Document()
        d.add_paragraph("前")
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "表"
        d.add_paragraph("后")
        kinds = ["table" if hasattr(b, "rows") else b.text for b in _iter_docx_blocks(d)]
        self.assertEqual(kinds, ["前", "table", "后"])

    def test_discover_prefers_import_dir(self):
        """发现逻辑：导入目录有文件就不看资产层（且返回 from_imports=True）。"""
        from vllm_kb.config import SourceCfg
        from vllm_kb.sources import _discover_source_files, WordSource

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            src = WordSource(SourceCfg(id="c", type="word", path="data/imports/word"),
                             project_root=root)
            (root / "data" / "imports" / "word").mkdir(parents=True)
            (root / "data" / "assets" / "word").mkdir(parents=True)
            (root / "data" / "imports" / "word" / "a.docx").write_bytes(b"x")
            (root / "data" / "assets" / "word" / "b.docx").write_bytes(b"y")
            files, fallback = _discover_source_files(
                src, ("*.docx",), src.import_dir, src._assets_dir())
            self.assertFalse(fallback)
            self.assertEqual([p.name for p, _, _ in files], ["a.docx"])
            self.assertTrue(all(fi for _, fi, _ in files))
            # 导入目录空 → 回退
            (root / "data" / "imports" / "word" / "a.docx").unlink()
            files, fallback = _discover_source_files(
                src, ("*.docx",), src.import_dir, src._assets_dir())
            self.assertTrue(fallback)
            self.assertEqual([p.name for p, _, _ in files], ["b.docx"])
            self.assertFalse(any(fi for _, fi, _ in files))
            # 导入目录不存在、资产层也不存在 → 空
            shutil.rmtree(root / "data" / "assets")
            self.assertEqual(
                _discover_source_files(src, ("*.docx",), src.import_dir, src._assets_dir()),
                ([], False))


class TestWordEmbeddedImages(_WordCase):
    """内嵌图片：不透明占位 + 内容寻址资产 + asset_registry 注册 + 正文不含路径。"""

    def _write_img(self, name="图片案例.docx", **kw) -> Path:
        p = self.import_dir / name
        make_docx_with_images(p, **kw)
        return p

    def test_placeholder_and_asset(self):
        """图片 → 正文 `[图片]` 占位（无路径），图片落 assets/images（内容寻址命名）。"""
        self._write_img()
        docs = self._docs()
        self.assertEqual(len(docs), 1)
        d = docs[0]
        self.assertIn("[图片]", d.body)
        self.assertEqual(d.extra["quality"]["images"], 1)
        # 资产以 sha 命名：不含 Word 文件名（包内 media 名在文档间必然重名）
        imgs = sorted((self.root / "data" / "assets" / "images").glob("*.png"))
        self.assertEqual(len(imgs), 1)
        self.assertRegex(imgs[0].name, r"^img_[0-9a-f]{16}\.png$")
        # evidence：只有 asset_id/sha256，无 source_ref（不记任何路径）
        ev = d.extra["evidence"][0]
        self.assertEqual(ev["kind"], "embedded")
        self.assertNotIn("source_ref", ev)
        self.assertEqual(ev["asset_id"], imgs[0].stem.split("_", 1)[1])
        # 正文/整条 extra 不含路径形态
        blob = json.dumps(d.extra, ensure_ascii=False)
        for bad in ("imports", "assets/", "parsed/", ".docx", "media"):
            self.assertNotIn(bad, blob, f"extra 泄漏 {bad!r}")

    def test_asset_registry_registration(self):
        """图片注册进 asset_registry（source_type=image）→ 审核台可反查路径并预览。"""
        from vllm_kb.review import list_assets

        os.environ["VLLM_KB_DATA_ROOT"] = str(self.root / "data")
        self._write_img()
        src = self._src(app_cfg=AppConfig.model_validate({}))
        src.pull()
        src.canonicalize()
        assets = list_assets(self.root / "data" / "review.sqlite3")
        by_type = {}
        for reg in assets.values():
            by_type.setdefault(reg["source_type"], []).append(reg["rel_path"])
        self.assertEqual(len(by_type.get("image", [])), 1)
        self.assertTrue(by_type["image"][0].startswith("assets/images/img_"))
        # word 原件与图片同批注册
        self.assertEqual(len(by_type.get("doc_word", [])), 1)

    def test_alt_text_in_placeholder(self):
        """`wp:docPr/@descr` → `[图片:alt]`（alt 是文档内容，随正文进 FTS）。"""
        self._write_img(alt="拓扑图")
        body = self._docs()[0].body
        self.assertIn("[图片:拓扑图]", body)

    def test_alt_text_normalized_and_truncated(self):
        """alt 压掉换行、限长（畸形 alt 能到几 KB，不该整段进正文）。"""
        from vllm_kb.sources import _word_alt_text

        self.assertEqual(_word_alt_text("a\n\n b\t c "), "a b c")
        self.assertEqual(len(_word_alt_text("x" * 500)), 120)

    def test_identical_images_share_one_asset(self):
        """同图重复出现 → 只落一份资产（内容寻址），但占位符各留一个。"""
        self._write_img(n=2, colors=[(7, 7, 7), (7, 7, 7)])
        d = self._docs()[0]
        self.assertEqual(d.body.count("[图片]"), 2)
        self.assertEqual(d.extra["quality"]["images"], 2)
        self.assertEqual(len(list((self.root / "data" / "assets" / "images").glob("*.png"))), 1)

    def test_distinct_images_two_assets(self):
        self._write_img(n=2, colors=[(1, 2, 3), (4, 5, 6)])
        self._docs()
        self.assertEqual(len(list((self.root / "data" / "assets" / "images").glob("*.png"))), 2)

    def test_orphan_image_registered_without_placeholder(self):
        """孤儿图（包内有关系、正文未引用）：注册资产但不插占位（没有正文落点）。"""
        self._write_img(orphan=True)
        d = self._docs()[0]
        self.assertEqual(d.extra["quality"]["images"], 1)               # 正文 1 张
        self.assertEqual(d.extra["quality"]["images_unreferenced"], 1)  # 孤儿 1 张
        self.assertEqual(d.body.count("[图片]"), 1)
        # 两张图都落了资产（孤儿图仍可被审核台预览）
        self.assertEqual(len(list((self.root / "data" / "assets" / "images").glob("*.png"))), 2)

    def test_image_only_document_not_skipped(self):
        """纯图片文档不再被整篇跳过（回归：以前 `无正文 → 跳过`，图里的信息完全丢失）。"""
        self._write_img(only_image=True)
        docs = self._docs()
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].body, "[图片]")
        self.assertEqual(docs[0].title, "图片案例")   # 无标题/正文 → 回退文件名

    def test_image_in_table_placeholder_after_table(self):
        """表格里的图片：占位符排在表格之后（塞进单元格会破 Markdown 表格）。"""
        import io

        import docx

        p = self.import_dir / "表内图.docx"
        d = docx.Document()
        d.add_heading("表内图", level=1)
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "错误码"
        t.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(make_png()))
        d.save(str(p))
        body = self._docs()[0].body
        self.assertIn("| 错误码 |", body)
        self.assertIn("[图片]", body)
        self.assertLess(body.index("| 错误码 |"), body.index("[图片]"))
        # 合并单元格不会让同一段落被重复枚举 → 只有 1 个占位
        self.assertEqual(body.count("[图片]"), 1)

    def test_external_link_image_placeholder_only(self):
        """外链图（`r:link`，内容不在包里）：照常占位，但不登记资产、不记链接（不泄漏路径）。"""
        import docx
        from docx.oxml import parse_xml

        p = self.import_dir / "外链图.docx"
        d = docx.Document()
        d.add_heading("外链图", level=1)
        par = d.add_paragraph()
        par._p.append(parse_xml(
            '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<v:shape><v:imagedata r:id="rIdExternal" r:link="rIdExternal"/></v:shape></w:pict>'))
        d.save(str(p))
        docs = self._docs()
        self.assertEqual(len(docs), 1)
        self.assertIn("[图片]", docs[0].body)
        self.assertEqual(docs[0].extra["quality"]["images"], 1)
        self.assertEqual(docs[0].extra["quality"]["images_unresolved"], 1)
        self.assertEqual(list((self.root / "data" / "assets" / "images").glob("*")), [])
        self.assertNotIn("rIdExternal", json.dumps(docs[0].extra, ensure_ascii=False))


class TestWordImageOoxml(unittest.TestCase):
    """OOXML 图片遍历的边界（手搓 XML：真实 Word 的兼容/旧版形态很难用 python-docx 造）。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.p = self.root / "t.docx"
        make_docx_with_images(self.p, n=1, heading=False)

    def tearDown(self):
        self.tmp.cleanup()

    def test_vml_imagedata_found(self):
        """VML（`v:imagedata`，兼容模式/旧版粘贴）也要认出来——只扫 a:blip 会漏图。"""
        import docx
        from docx.oxml import parse_xml

        from vllm_kb.sources import _word_images_in

        rid = _docx_rId(self.p)
        d = docx.Document(str(self.p))
        # 摘掉原 drawing，换成 VML 形态引用同一 rId
        drawing = d.paragraphs[1]._p.xpath(".//w:drawing")[0]
        drawing.getparent().remove(drawing)
        d.paragraphs[1]._p.append(parse_xml(
            '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<v:shape><v:imagedata r:id="{rid}"/></v:shape></w:pict>'))
        got = _word_images_in(d.paragraphs[1]._p)
        self.assertEqual([r for r, _ in got], [rid])

    def test_alternate_content_fallback_not_double_counted(self):
        """`mc:AlternateContent` 的 Choice 与 Fallback 引用同一 rId → 只能算一张图。

        全遍历会把一张图数成两张：正文出现两个占位、资产重复登记、OCR 白跑一次。
        """
        import docx
        from docx.oxml import parse_xml

        from vllm_kb.sources import _word_images_in

        rid = _docx_rId(self.p)
        d = docx.Document(str(self.p))
        drawing = d.paragraphs[1]._p.xpath(".//w:drawing")[0]
        drawing.getparent().remove(drawing)
        d.paragraphs[1]._p.append(parse_xml(
            '<mc:AlternateContent '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:v="urn:schemas-microsoft-com:vml">'
            '<mc:Choice Requires="wps"><w:drawing><wp:inline><a:graphic>'
            f'<a:blip r:embed="{rid}"/></a:graphic></wp:inline></w:drawing></mc:Choice>'
            f'<mc:Fallback><w:pict><v:shape><v:imagedata r:id="{rid}"/>'
            '</v:shape></w:pict></mc:Fallback></mc:AlternateContent>'))
        got = _word_images_in(d.paragraphs[1]._p)
        self.assertEqual([r for r, _ in got], [rid], "Fallback 分支被重复计数")

    def test_comment_nodes_ignored(self):
        """XML 注释/PI 不是元素（`.tag` 是函数），遍历必须跳过而不是崩。"""
        import docx
        from docx.oxml import parse_xml

        from vllm_kb.sources import _word_images_in

        d = docx.Document(str(self.p))
        par = d.add_paragraph()          # 新段落：本身没有图片
        par._p.append(parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<!-- 注释 --><w:t>x</w:t></w:r>'))
        self.assertEqual(_word_images_in(par._p), [])


class _WordOcrCase(_WordCase):
    """OCR 用例：需要 app_cfg（OCR 配置来自 image source）+ 数据根重定向。"""

    def setUp(self):
        super().setUp()
        os.environ["VLLM_KB_DATA_ROOT"] = str(self.root / "data")

    def _app(self, min_conf: float = 0.6) -> AppConfig:
        return AppConfig.model_validate({
            "embedding": {"provider": "echo", "dimensions": 8},
            "storage": {"vector_backend": "python", "review_path": "data/review.sqlite3"},
            "sources": [
                {"id": "cases", "type": "word", "path": "data/imports/word", "enabled": True},
                {"id": "images", "type": "image", "ocr_provider": "api",
                 "ocr_api_base": "http://ocr.local:8000",
                 "ocr_min_confidence": min_conf, "enabled": True},
            ],
        })

    def _docs_ocr(self, min_conf: float = 0.6):
        if not any(self.import_dir.rglob("*.docx")):
            make_docx_with_images(self.import_dir / "图片案例.docx")
        src = self._src(app_cfg=self._app(min_conf))
        src.pull()
        return src.canonicalize()

    def _img_docx(self) -> Path:
        """给直接调 `_materialize_images` 的用例一个真实 docx（它要打开包读 rels）。"""
        p = self.import_dir / "直接调用.docx"
        if not p.exists():
            make_docx_with_images(p)
        return p


class TestWordImageOcr(_WordOcrCase):
    """内嵌图片 OCR：高置信注入正文、低置信只留线索、阈值不被缓存冻结。"""

    def _mock(self, text: str, conf: float):
        from vllm_kb.ocr import OcrResult

        return mock.patch("vllm_kb.ocr.ocr_image_detail",
                          return_value=OcrResult(text=text, confidence=conf,
                                                 confidence_source="engine", provider="api"))

    def test_high_confidence_text_injected(self):
        """高置信 OCR 文本注入占位符之后 → 随正文进 FTS + 向量。"""
        with self._mock("halMemCreate failed drvRetCode=6", 0.92):
            d = self._docs_ocr()[0]
        self.assertIn("halMemCreate failed drvRetCode=6", d.body)
        self.assertLess(d.body.index("[图片]"), d.body.index("halMemCreate"))
        ev = d.extra["evidence"][0]
        self.assertEqual(ev["kind"], "embedded")
        self.assertTrue(ev["ocr"]["text_included"])
        self.assertEqual(ev["ocr"]["confidence"], 0.92)

    def test_low_confidence_not_injected_but_queued(self):
        """低置信/自报异常：不注入正文，只留签名线索 + evidence（→ 审核队列据此入队）。"""
        with self._mock("maybe 561000", 0.30):
            d = self._docs_ocr()[0]
        self.assertNotIn("maybe 561000", d.body)
        self.assertIn("[图片]", d.body)
        ocr = d.extra["evidence"][0]["ocr"]
        self.assertFalse(ocr["text_included"])
        self.assertEqual(ocr["confidence"], 0.3)

    def test_ocr_text_not_frozen_in_extract_cache(self):
        """**关键不变量**：OCR 文本不进 extract 缓存 → 调阈值只重判定、不重解析。

        缓存里若冻结了 OCR 结果，会出现两个坑：① 调 `ocr_min_confidence` 不生效（要清缓存）；
        ② 首次构建时 OCR 服务不可用 → 空文本被永久冻结，服务恢复后也永远补不回来。
        """
        with self._mock("halMemCreate failed drvRetCode=6", 0.70):
            first = self._docs_ocr(min_conf=0.9)[0]      # 0.70 < 0.9 → 不注入
        self.assertNotIn("halMemCreate", first.body)
        cache = next((self.root / "data" / "parsed" / "word").glob("*.extract.json"))
        cached = json.loads(cache.read_text(encoding="utf-8"))
        self.assertNotIn("halMemCreate", cached["body"], "OCR 文本被冻结进 extract 缓存")
        self.assertIn("[图片]", cached["body"])
        # 同一份 extract 缓存（sha 未变）+ 放宽阈值 → 立刻注入，**不需要清缓存**
        with self._mock("halMemCreate failed drvRetCode=6", 0.70):
            second = self._docs_ocr(min_conf=0.5)[0]
        self.assertIn("halMemCreate failed drvRetCode=6", second.body)
        self.assertEqual(cache.read_text(encoding="utf-8"),
                         json.dumps(cached, ensure_ascii=False), "缓存被重写了（应命中）")

    def test_ocr_unavailable_degrades_then_recovers(self):
        """OCR 不可用 → 导入不受阻（只占位）；服务恢复后同一缓存即注入（不冻结空结果）。"""
        from vllm_kb.ocr import OcrApiError

        with mock.patch("vllm_kb.ocr.ocr_image_detail", side_effect=OcrApiError("svc down")):
            d = self._docs_ocr()[0]
        self.assertEqual(d.body.count("[图片]"), 1)
        self.assertIsNone(d.extra["evidence"][0]["ocr"])
        with self._mock("recovered text 561000", 0.9):
            d2 = self._docs_ocr()[0]
        self.assertIn("recovered text 561000", d2.body)

    def test_ocr_cache_reused_across_runs(self):
        """同一张图 OCR 结果按 sha 缓存（ocr.json）→ 第二次构建不再调 OCR。"""
        with self._mock("cached text", 0.9) as m:
            self._docs_ocr()
            calls_first = m.call_count
        with self._mock("cached text", 0.9) as m2:
            self._docs_ocr()
            calls_second = m2.call_count
        self.assertEqual(calls_first, 1)
        self.assertEqual(calls_second, 0, "第二次构建仍调了 OCR（幂等缓存失效）")

    def test_non_raster_not_sent_to_ocr(self):
        """矢量/多页格式（emf/wmf/svg/tiff）注册资产但不送 OCR（引擎不收，白花调用）。"""
        entries = [{"rid": "rId1", "sha256": "a" * 64, "ext": "emf",
                    "alt": "", "in_body": True}]
        src = self._src(app_cfg=self._app())
        with mock.patch.object(WordSource, "_write_image") as w, \
                mock.patch("vllm_kb.ocr.ocr_image_detail") as ocr_m:
            w.return_value = self.root / "data" / "assets" / "images" / "img_x.emf"
            evidence, injects, registered = src._materialize_images(
                self._img_docx(), {"images": entries})
        self.assertEqual(injects, [""])
        ocr_m.assert_not_called()
        self.assertEqual(evidence[0]["sha256"], "a" * 64)
        self.assertEqual(len(registered), 1)

    def test_orphan_image_not_ocrd(self):
        """孤儿图不做 OCR（没有正文落点，文本无处可去）——只注册资产。"""
        entries = [{"rid": "rId1", "sha256": "b" * 64, "ext": "png",
                    "alt": "", "in_body": False}]
        src = self._src(app_cfg=self._app())
        with mock.patch.object(WordSource, "_write_image") as w, \
                mock.patch("vllm_kb.ocr.ocr_image_detail") as ocr_m:
            w.return_value = self.root / "data" / "assets" / "images" / "img_y.png"
            evidence, injects, registered = src._materialize_images(
                self._img_docx(), {"images": entries})
        self.assertEqual(injects, [])
        ocr_m.assert_not_called()
        self.assertEqual(len(registered), 1)
        self.assertIsNone(evidence[0]["ocr"])


class TestWordExtractCacheSchema(_WordCase):
    """extract 缓存的 schema 位：提取逻辑升级必须让旧缓存失效。"""

    def test_stale_schema_cache_invalidated(self):
        """旧 schema 缓存（sha 仍匹配）必须被重新解析——否则图片**静默不生效**。"""
        p = self.import_dir / "图片案例.docx"
        make_docx_with_images(p, n=1)
        src = self._src()
        src.pull()
        docs = src.canonicalize()
        self.assertIn("[图片]", docs[0].body)
        cache = next((self.root / "data" / "parsed" / "word").glob("*.extract.json"))
        # 伪造 step1（无图片清单）的旧缓存：sha 一致、schema 缺失、body 无占位
        stale = json.loads(cache.read_text(encoding="utf-8"))
        stale.pop("schema")
        stale["body"] = "# 图片案例\n\n下图是拓扑："
        stale.pop("images")
        stale.pop("img_offsets")
        cache.write_text(json.dumps(stale, ensure_ascii=False), encoding="utf-8")
        # 重跑：schema 不符 → 重新解析 → 图片回来了
        again = self._src().canonicalize()
        self.assertIn("[图片]", again[0].body)
        self.assertEqual(again[0].extra["quality"]["images"], 1)
        self.assertEqual(json.loads(cache.read_text(encoding="utf-8")).get("schema"), 2)

    def test_matching_schema_cache_hits(self):
        """schema 一致 → 命中缓存（不重写文件，mtime 不变）。"""
        p = self.import_dir / "图片案例.docx"
        make_docx_with_images(p, n=1)
        src = self._src()
        src.pull()
        src.canonicalize()
        cache = next((self.root / "data" / "parsed" / "word").glob("*.extract.json"))
        before = cache.stat().st_mtime_ns
        docs = self._src().canonicalize()
        self.assertEqual(cache.stat().st_mtime_ns, before, "缓存被重写（应命中）")
        self.assertIn("[图片]", docs[0].body)


def _add_hyperlink(par, text: str, url: str = "", *, anchor: str = "") -> None:
    """给段落加超链接（python-docx 没有 add_hyperlink，手搓）。

    url 非空 → 建**外部关系**（`w:hyperlink/@r:id`）；anchor 非空 → 内部书签（`w:anchor`，无 URL）。
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    h = OxmlElement("w:hyperlink")
    if anchor:
        h.set(qn("w:anchor"), anchor)
    else:
        h.set(qn("r:id"), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    par._p.append(h)


def _add_style_with_level(d, name: str, level: int, *, style_id: str = "") -> object:
    """建一个带 `w:outlineLvl` 的段落样式（模拟用户自定义/本地化的标题样式）。

    `style_id` 可覆盖——用来隔离"靠样式名匹配"与"靠大纲级别匹配"两条路径。
    """
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn

    st = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if style_id:
        st.element.set(qn("w:styleId"), style_id)
    st.element.append(parse_xml(
        f'<w:pPr {nsdecls("w")}><w:outlineLvl w:val="{level - 1}"/></w:pPr>'))
    return st


class TestWordHyperlinks(_WordCase):
    """超链接 → `[text](url)`：`Paragraph.text` 会并进链接文字但**丢掉地址**。"""

    def _doc_with(self, build) -> Path:
        import docx

        p = self.import_dir / "链接.docx"
        d = docx.Document()
        d.add_heading("链接案例", level=1)
        build(d)
        d.save(str(p))
        return p

    def test_external_hyperlink_rendered_with_url(self):
        def build(d):
            par = d.add_paragraph("见 ")
            _add_hyperlink(par, "昇腾社区文档", "https://www.hiascend.com/document")
            par.add_run(" 的说明")

        self._doc_with(build)
        body = self._docs()[0].body
        self.assertIn("[昇腾社区文档](https://www.hiascend.com/document)", body)
        self.assertIn("见 ", body)
        self.assertIn(" 的说明", body)

    def test_internal_anchor_text_only(self):
        """内部书签（`w:anchor`）没有 URL → 只留文字，不能渲染成空链接 `[x]()`。"""
        self._doc_with(lambda d: _add_hyperlink(d.add_paragraph("跳转："), "见第 3 章",
                                                anchor="chap3"))
        body = self._docs()[0].body
        self.assertIn("见第 3 章", body)
        self.assertNotIn("见第 3 章]", body)
        self.assertNotIn("]()", body)

    def test_non_http_target_not_written_to_body(self):
        """**路径安全**：file:// / UNC 等非 http(s) 目标只留链接文字，地址不写进正文。"""
        self._doc_with(lambda d: _add_hyperlink(
            d.add_paragraph("附件："), "内部附件", "file:///D:/internal/secret/plan.docx"))
        body = self._docs()[0].body
        self.assertIn("内部附件", body)
        self.assertNotIn("secret", body)
        self.assertNotIn("file://", body)
        self.assertNotIn("D:/", body)

    def test_hyperlink_in_table_cell(self):
        """表格单元格里的超链接也要渲染（否则表内链接地址同样丢失）。"""
        def build(d):
            t = d.add_table(rows=1, cols=2)
            t.cell(0, 0).text = "参考"
            _add_hyperlink(t.cell(0, 1).paragraphs[0], "手册", "https://example.com/manual")

        self._doc_with(build)
        body = self._docs()[0].body
        self.assertIn("[手册](https://example.com/manual)", body)

    def test_hyperlink_inside_inserted_run(self):
        """`w:ins`（修订插入）是透明容器：里面的文字（含超链接）也是正文。"""
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        def build(d):
            par = d.add_paragraph("开始 ")
            par._p.append(parse_xml(
                f'<w:ins {nsdecls("w")} w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z">'
                '<w:r><w:t>插入的正文</w:t></w:r></w:ins>'))

        self._doc_with(build)
        self.assertIn("插入的正文", self._docs()[0].body)


class TestWordHeadingStyles(unittest.TestCase):
    """标题识别三重信号：样式名/ID 正则 → 名字表 → `w:outlineLvl`。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_chinese_style_name_matched_by_name(self):
        """中文界面样式名「标题 3」——**故意把 styleId 改成中性值**，隔离出"靠名字匹配"。"""
        import docx

        from vllm_kb.sources import _word_heading_level

        d = docx.Document()
        st = d.styles.add_style("标题 3", 1)          # 1 = PARAGRAPH
        st.element.set(__import__("docx").oxml.ns.qn("w:styleId"), "Neutral3")
        par = d.add_paragraph("三级标题", style="标题 3")
        self.assertEqual(par.style.style_id, "Neutral3")
        self.assertEqual(_word_heading_level(par), 3)

    def test_custom_style_name_falls_back_to_outline_level(self):
        """自定义样式名（"我的章节"）靠 `w:outlineLvl` 识别——名字正则救不了。"""
        import docx

        from vllm_kb.sources import _word_heading_level

        d = docx.Document()
        _add_style_with_level(d, "我的章节", 2)
        par = d.add_paragraph("自定义章节", style="我的章节")
        self.assertEqual(_word_heading_level(par), 2)

    def test_paragraph_direct_outline_level(self):
        """段落直接格式上的 `w:outlineLvl`（样式名完全中性）也要认。"""
        import docx
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        from vllm_kb.sources import _word_heading_level

        d = docx.Document()
        par = d.add_paragraph("直接格式标题")
        par._p.get_or_add_pPr().append(
            parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
        self.assertEqual(_word_heading_level(par), 1)

    def test_outline_level_body_text_not_heading(self):
        """`w:outlineLvl val=9`（正文级）不是标题。"""
        import docx
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        from vllm_kb.sources import _word_heading_level

        d = docx.Document()
        par = d.add_paragraph("正文")
        par._p.get_or_add_pPr().append(
            parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="9"/>'))
        self.assertEqual(_word_heading_level(par), 0)

    def test_subtitle_is_level_2(self):
        import docx

        from vllm_kb.sources import _word_heading_level

        d = docx.Document()
        # 默认模板自带 Subtitle 样式（add_style 会报"已存在"）
        self.assertEqual(_word_heading_level(d.add_paragraph("副标题", style="Subtitle")), 2)

    def test_list_style_not_mistaken_for_heading(self):
        """列表样式不能被大纲级别兜底误判成标题（否则列表项会变成章节）。"""
        import docx

        from vllm_kb.sources import _word_heading_level, _word_list_prefix

        d = docx.Document()
        par = d.add_paragraph("列表项", style="List Bullet")
        self.assertEqual(_word_heading_level(par), 0)
        self.assertEqual(_word_list_prefix(par), "- ")

    def test_block_level_sdt_unwrapped(self):
        """块级内容控件（`w:sdt`）包住的段落/表格要下潜取出，否则整块内容丢失。"""
        import docx
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        from vllm_kb.sources import _iter_docx_blocks

        d = docx.Document()
        d.add_paragraph("前")
        # 必须插在 w:sectPr 之前：body.append 会落到节属性之后，顺序断言就失去意义
        d.element.body.sectPr.addprevious(parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtContent>'
            '<w:p><w:r><w:t>控件里的段落</w:t></w:r></w:p>'
            '</w:sdtContent></w:sdt>'))
        d.add_paragraph("后")
        texts = [b.text if not hasattr(b, "rows") else "table" for b in _iter_docx_blocks(d)]
        self.assertEqual(texts, ["前", "控件里的段落", "后"])


class TestWordBadFiles(_WordCase):
    """旧格式/加密/损坏文件：跳过原因要**可操作**（否则用户不知道下一步做什么）。"""

    def test_legacy_files_get_targeted_hint(self):
        """目录里只有 .doc → 明确提示另存为 .docx，而不是干巴巴一句"无 Word 文件"。"""
        from vllm_kb.sources import _CFB_MAGIC

        (self.import_dir / "旧案例.doc").write_bytes(_CFB_MAGIC + b"\x00" * 64)
        (self.import_dir / "说明.rtf").write_text("{\\rtf1}", encoding="utf-8")
        src = self._src()
        hint = src._legacy_hint()
        self.assertIn("旧案例.doc", hint)
        self.assertIn("说明.rtf", hint)
        self.assertIn("另存为 .docx", hint)
        # 直接调 canonicalize：`_docs()` 会自动补一个默认 docx，掩盖"无 Word 文件"的场景
        src.pull()
        self.assertEqual(src.canonicalize(), [])

    def test_no_legacy_hint_when_dir_empty(self):
        self.assertEqual(self._src()._legacy_hint(), "")

    def test_cfb_docx_reports_encrypted_or_legacy(self):
        """CFB 容器伪装成 .docx（= 加密/受保护 docx，或改了扩展名的 .doc）要说明白。"""
        from vllm_kb.sources import _CFB_MAGIC, _word_bad_file_reason

        p = self.import_dir / "加密.docx"
        p.write_bytes(_CFB_MAGIC + b"\x00" * 64)
        reason = _word_bad_file_reason(p)
        self.assertIn("加密", reason)
        self.assertIn("另存为", reason)
        self.assertIsNone(self._src()._extract_docx(p))

    def test_non_zip_reports_not_ooxml(self):
        from vllm_kb.sources import _word_bad_file_reason

        p = self.import_dir / "伪装.docx"
        p.write_bytes(b"just plain text, not a zip")
        self.assertIn("不是 OOXML", _word_bad_file_reason(p))

    def test_broken_zip_reports_structure_damage(self):
        from vllm_kb.sources import _word_bad_file_reason

        p = self.import_dir / "损坏.docx"
        p.write_bytes(b"PK\x03\x04" + b"\x00" * 64)     # zip 魔数但内容无效
        self.assertIn("OOXML 结构损坏", _word_bad_file_reason(p))
        self.assertIsNone(self._src()._extract_docx(p))

    def test_unreadable_file_reports_io(self):
        from vllm_kb.sources import _word_bad_file_reason

        self.assertIn("无法读取", _word_bad_file_reason(self.import_dir / "不存在.docx"))


if __name__ == "__main__":
    unittest.main()
