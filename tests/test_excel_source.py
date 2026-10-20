"""Excel 来源（schema-free）测试：多 sheet/变列解析、脱敏、实体复用入图。"""
import json
import tempfile
import unittest
from pathlib import Path

from vllm_kb.config import AppConfig, SourceCfg
from vllm_kb.sources import ExcelSource


def make_xlsx(path: Path) -> None:
    """生成含 2 个 sheet、列数不同、含空行的测试 xlsx。

    - Sheet1：表头（问题/现象/结论）+ 3 数据行，含 IP 与内部路径 cell；
    - Sheet2：2 列单数据行（无表头结构——行数 <3，验证裸拼回退）。
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["问题", "现象", "结论"])
    ws.append(["HCCL 超时", "节点 10.0.0.5 通信超时", "检查 /var/log/npu/ 日志"])
    ws.append(["内存不足", "OOM", "路径 /home/user/logs/oom.log 有记录"])
    ws.append([None, None, None])  # 空行应跳过
    ws.append(["回环测试", "127.0.0.1 正常", ""])
    ws2 = wb.create_sheet("任意名")
    ws2.append(["错误码 561000", "aclnn 失败"])  # 2 列——列数不同也应解析
    wb.save(str(path))


class TestExcelSource(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        imports = self.root / "data" / "imports" / "engineer"
        imports.mkdir(parents=True)
        make_xlsx(imports / "问题定位记录.xlsx")
        self.cfg = SourceCfg(id="engineer", type="excel",
                             path="data/imports/engineer/问题定位记录.xlsx", enabled=True)
        self.src = ExcelSource(self.cfg, project_root=self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_pull_copies_to_assets(self):
        n = self.src.pull()
        self.assertEqual(n, 1)
        self.assertTrue((self.root / "data" / "assets" / "excel" / "问题定位记录.xlsx").exists())
        # 幂等
        self.assertEqual(self.src.pull(), 0)

    def test_canonicalize_one_doc_per_file(self):
        """**一个文件一篇**：多 sheet/变列/空行都进同一篇 body，sheet 名作 Markdown 标题。

        身份 = `excel:<文件名>`（不再是位置性的 `文件:sheet:行`）——插行/改行不会让身份平移。
        """
        self.src.pull()
        docs = self.src.canonicalize()
        self.assertEqual(len(docs), 1)
        doc = docs[0]
        self.assertEqual(doc.source_id, "excel:问题定位记录")
        self.assertEqual(doc.title, "问题定位记录")
        self.assertEqual(doc.source_type, "doc_excel")
        self.assertEqual(doc.extra["verification"], "unverified")
        # 每个 sheet 一个小节标题
        self.assertIn("## Sheet1", doc.body)
        self.assertIn("## 任意名", doc.body)
        # 表头键值对：数据行自包含列语义
        self.assertIn("[表头] 问题 现象 结论", doc.body)
        self.assertIn("问题=内存不足 | 现象=OOM", doc.body)
        self.assertIn("问题=HCCL 超时 | 现象=节点 10.0.0.5 通信超时", doc.body)
        self.assertIn("561000", doc.body)      # 变列（Sheet2 2 列）也解析
        self.assertIn("aclnn", doc.body)
        # 空行不进正文（Sheet1 第 4 行）
        self.assertEqual(doc.body.count("\n\n## "), 1)   # 两个 sheet 之间只有一个分隔
        # 计数元数据（整体评估用；quality 出口整包透传）
        self.assertEqual(doc.extra["quality"]["sheets"], 2)
        self.assertEqual(doc.extra["quality"]["rows"], 5)   # 4 + 1 个有效行
        self.assertEqual(doc.extra["quality"]["sheets_skipped"], [])
        self.assertEqual(doc.extra["quality"]["source_mode"], "imports")

    def test_chunks_are_still_per_row_with_sheet_section(self):
        """按文件聚合**不损失检索粒度**：每个非空行仍是独立 chunk，且带 sheet 作 section。"""
        from vllm_kb.chunking import chunk_doc

        self.src.pull()
        doc = self.src.canonicalize()[0]
        chunks = chunk_doc(doc)
        texts = [c.text for c in chunks]
        self.assertEqual(len(chunks), 5, texts)
        self.assertEqual({c.section for c in chunks}, {"Sheet1", "任意名"})
        hit = next(c for c in chunks if "HCCL 超时" in c.text)
        self.assertEqual(hit.section, "Sheet1")
        self.assertTrue(hit.text.startswith("【Sheet1】"), hit.text)   # 标题注入正文
        self.assertIn("问题=HCCL 超时", hit.text)   # 行 chunk 自包含列语义
        # 每行一个 chunk（不跨行合并成一大块）
        self.assertTrue(all("561000" not in c.text or c.section == "任意名" for c in chunks))

    def test_file_identity_survives_row_insert(self):
        """**核心动机**：中间插一行不会改变文档身份（逐行方案下会整体平移）。"""
        import openpyxl

        self.src.pull()
        before = self.src.canonicalize()[0].source_id
        p = self.root / "data" / "imports" / "engineer" / "问题定位记录.xlsx"
        wb = openpyxl.load_workbook(str(p))
        wb["Sheet1"].insert_rows(2)          # 在表头与第二行之间插一行
        wb["Sheet1"].cell(row=2, column=1, value="新增问题")
        wb.save(str(p))
        after = self.src.canonicalize()[0].source_id
        self.assertEqual(before, after)
        self.assertEqual(after, "excel:问题定位记录")

    def test_tags_from_filename_and_sheets(self):
        """标签是**整体评估**：文件名 + sheet 名走与 md/word 同一套 extract_tags。"""
        from vllm_kb.tagging import TagRegistry

        self.src.pull()
        doc = self.src.canonicalize()[0]
        self.assertIsInstance(doc.tags, list)          # 无词典命中时为空列表
        self.assertIn("tag_candidates", doc.extra)     # 候选走审核队列
        self.assertEqual(TagRegistry().entries, [])

    def test_doc_asset_registered(self):
        """§1.6：canonicalize 也把 excel 原件落资产层并注册（--skip-pull 也要有预览链）。"""
        import os

        from vllm_kb.config import AppConfig
        from vllm_kb.review import list_assets

        os.environ["VLLM_KB_DATA_ROOT"] = str(self.root / "data")
        try:
            cfg = AppConfig.model_validate({})
            src = ExcelSource(self.cfg, project_root=self.root, app_cfg=cfg)
            docs = src.canonicalize()          # 注意：**不调 pull()**
            self.assertEqual(len(docs), 1)
            assets = list_assets(self.root / "data" / "review.sqlite3")
            reg = assets.get(docs[0].extra["asset"]["asset_id"])
            self.assertIsNotNone(reg, assets)
            self.assertEqual(reg["source_type"], "doc_excel")
            self.assertEqual(reg["rel_path"], "assets/excel/问题定位记录.xlsx")
            self.assertTrue((self.root / "data" / reg["rel_path"]).is_file())
        finally:
            os.environ.pop("VLLM_KB_DATA_ROOT", None)

    def test_empty_file_skipped(self):
        """所有 sheet 都为空 → 不产文档（但原件照常落资产层）。"""
        import openpyxl

        p = self.root / "data" / "imports" / "engineer" / "空表.xlsx"
        wb = openpyxl.Workbook()
        wb.save(str(p))
        cfg = SourceCfg(id="e2", type="excel", path=str(p), enabled=True)
        self.assertEqual(ExcelSource(cfg, project_root=self.root).canonicalize(), [])

    def test_fallback_to_assets_when_imports_missing(self):
        """imports 缺失 → 回退资产层副本，身份取族名，quality 标记来源模式。"""
        self.src.pull()
        (self.root / "data" / "imports" / "engineer" / "问题定位记录.xlsx").unlink()
        docs = self.src.canonicalize()
        self.assertEqual([d.source_id for d in docs], ["excel:问题定位记录"])
        self.assertEqual(docs[0].extra["quality"]["source_mode"], "assets_fallback")

    def test_raw_text_stored_sanitize_log_collected(self):
        """后置脱敏：body 原文入库（原文检索）；会被脱敏的 IP/路径落盘维护日志。"""
        self.src.pull()
        docs = self.src.canonicalize()
        body = "\n".join(d.body for d in docs)
        # 原文入库（不脱敏——出口脱敏由 serve_api 返回时做）
        self.assertIn("10.0.0.5", body)
        self.assertIn("/home/user/logs/oom.log", body)
        self.assertIn("/var/log/npu/", body)
        self.assertIn("127.0.0.1", body)

    def test_entities_reuse_existing_pipeline(self):
        """scheme-free：body 实体（错误码）复用现有提取——canonical 后入图走 MENTIONS。"""
        from vllm_kb.graph_rels import extract_doc_relations

        self.src.pull()
        docs = self.src.canonicalize()
        doc = docs[0]
        ex = extract_doc_relations(doc.source_id, "", 0, doc.source_type, doc.body)
        self.assertIn("561000", ex.mentions.get("error_code", set()))
        # 有 body 的文档都能被 canonical/ingest 接受
        self.assertTrue(doc.body)

    def test_sanitize_log_written(self):
        """被脱敏命中的 IP/路径落盘 data/sanitize_log.json（维护用，含被脱敏值）。"""
        import json
        import os

        from vllm_kb.config import AppConfig

        os.environ["VLLM_KB_DATA_ROOT"] = str(self.root / "data")
        try:
            cfg = AppConfig.model_validate({})
            src = ExcelSource(self.cfg, project_root=self.root, app_cfg=cfg)
            src.pull()
            src.canonicalize()
            log = self.root / "data" / "sanitize_log.json"
            self.assertTrue(log.exists())
            data = json.loads(log.read_text(encoding="utf-8"))
            self.assertIn("10.0.0.5", data["ips"])          # 被脱敏的 IP
            self.assertIn("/home/user/logs/oom.log", data["paths"])  # 被脱敏的路径
            # 保留项不进入日志
            self.assertNotIn("/var/log/npu/", data["paths"])
        finally:
            os.environ.pop("VLLM_KB_DATA_ROOT", None)

    def test_sanitize_config_from_app_cfg(self):
        """config.sanitize 控制**收集范围**（后置脱敏：body 始终原文入库）。"""
        import json
        import os

        from vllm_kb.config import AppConfig

        os.environ["VLLM_KB_DATA_ROOT"] = str(self.root / "data")
        try:
            cfg = AppConfig.model_validate({"sanitize": {
                "keep_paths": ["/home/user/logs"],  # 覆盖默认：业务日志目录保留（不收集）
                "keep_ips": ["10.0.0.5"],           # 覆盖默认：该私有 IP 保留（不收集）
            }})
            src = ExcelSource(self.cfg, project_root=self.root, app_cfg=cfg)
            src.pull()
            docs = src.canonicalize()
            body = "\n".join(d.body for d in docs)
            # 原文始终入库（出口脱敏由 serve_api 做）
            self.assertIn("10.0.0.5", body)
            self.assertIn("/home/user/logs/oom.log", body)
            # 维护日志按配置收集：keep 的不收集；覆盖后默认路径不再保留 → 被收集
            log = self.root / "data" / "sanitize_log.json"
            data = json.loads(log.read_text(encoding="utf-8"))
            self.assertNotIn("10.0.0.5", data["ips"])
            self.assertNotIn("/home/user/logs/oom.log", data["paths"])
            self.assertIn("/var/log/npu/", data["paths"])
        finally:
            os.environ.pop("VLLM_KB_DATA_ROOT", None)


def _save_wb(wb, path: Path) -> None:
    wb.save(str(path))


def _make_multi_sheet_xlsx(path: Path) -> None:
    """模拟真实多页业务 Excel：封面/前言/如何获取EventID/故障描述表/运维级别定义标准。"""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "封面"
    ws.append(["A3 健康管理故障定义"])
    ws.append(["版权所有 © 华为技术有限公司 2026。 保留一切权利。"])
    ws2 = wb.create_sheet("前言")
    ws2.append(["商标声明", "和其他华为商标均为华为技术有限公司的商标。"])
    ws3 = wb.create_sheet("如何获取EventID")
    ws3.append(["EventID 通过 BMC 命令行查询获取。"])
    ws4 = wb.create_sheet("故障描述表")
    ws4.append(["EventID", "故障名称", "处理建议"])
    ws4.append(["0x220001", "内存ECC可纠正错误", "观察运行状态"])
    ws4.append(["0x220002", "内存ECC不可纠正错误", "更换内存"])
    ws5 = wb.create_sheet("运维级别定义标准")
    ws5.append(["级别", "定义"])
    ws5.append(["严重", "影响业务运行"])
    _save_wb(wb, path)


class TestBoilerplateSheets(unittest.TestCase):
    """结构页过滤：封面/前言/目录/版权类 sheet 不进正文与标签（2026-10 业务实测）。

    业务多页 Excel 的首个 sheet 常是封面（"版权所有 ©…"）——不滤则文档 body 开头
    是版权样板，审核页标题观感受污染；sheet 名"封面/前言"也不该混进标签提取。
    """

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        imports = self.root / "data" / "imports" / "engineer"
        imports.mkdir(parents=True)
        _make_multi_sheet_xlsx(imports / "A3 健康管理故障定义.xlsx")
        self.cfg = SourceCfg(id="engineer", type="excel",
                             path="data/imports/engineer/A3 健康管理故障定义.xlsx",
                             enabled=True)
        self.src = ExcelSource(self.cfg, project_root=self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_boilerplate_sheets_filtered(self):
        self.src.pull()
        docs = self.src.canonicalize()
        self.assertEqual(len(docs), 1)
        doc = docs[0]
        # 封面/前言不进正文：body 不含版权/商标样板，开头是第一个内容 sheet
        self.assertNotIn("版权所有", doc.body)
        self.assertNotIn("商标声明", doc.body)
        self.assertTrue(doc.body.startswith("## 如何获取EventID"), doc.body[:80])
        # 保留的 sheet 正常入库（表头键值对）
        self.assertIn("## 故障描述表", doc.body)
        self.assertIn("EventID=0x220001 | 故障名称=内存ECC可纠正错误", doc.body)
        # 标签提取不含被滤 sheet 名
        self.assertNotIn("封面", doc.tags)
        self.assertNotIn("前言", doc.tags)
        # quality 可观测：滤了哪些
        q = doc.extra["quality"]
        self.assertEqual(q["sheets"], 3)
        self.assertEqual(q["sheets_skipped"], ["封面", "前言"])

    def test_skip_sheets_config_extends_filter(self):
        """config `skip_sheets`（正则数组）追加业务特有结构页名单。"""
        cfg = SourceCfg(id="engineer", type="excel",
                        path="data/imports/engineer/A3 健康管理故障定义.xlsx",
                        skip_sheets=["^运维级别"], enabled=True)
        src = ExcelSource(cfg, project_root=self.root)
        src.pull()
        doc = src.canonicalize()[0]
        self.assertNotIn("## 运维级别定义标准", doc.body)
        self.assertEqual(doc.extra["quality"]["sheets_skipped"],
                         ["封面", "前言", "运维级别定义标准"])

    def test_all_sheets_filtered_yields_no_doc(self):
        """全部 sheet 都是结构页 → 不产生文档（与全空 sheet 同径）。"""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "封面"
        ws.append(["版权所有 © 某公司"])
        p = self.root / "data" / "imports" / "engineer" / "只有封面.xlsx"
        _save_wb(wb, p)
        cfg = SourceCfg(id="engineer", type="excel",
                        path="data/imports/engineer/只有封面.xlsx", enabled=True)
        src = ExcelSource(cfg, project_root=self.root)
        src.pull()
        self.assertEqual(src.canonicalize(), [])


class TestHeaderKv(unittest.TestCase):
    """表头键值对 auto 判定边界：无表头结构回退裸拼（语义无损优先）。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        imports = self.root / "data" / "imports" / "engineer"
        imports.mkdir(parents=True)
        self.imports = imports

    def tearDown(self):
        self.tmp.cleanup()

    def _doc_body(self, name: str) -> str:
        cfg = SourceCfg(id="engineer", type="excel",
                        path=f"data/imports/engineer/{name}", enabled=True)
        src = ExcelSource(cfg, project_root=self.root)
        src.pull()
        docs = src.canonicalize()
        self.assertEqual(len(docs), 1)
        return docs[0].body

    def test_few_rows_no_kv(self):
        """非空行 <3：说明页形态 → 裸拼（首行不是表头）。"""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "说明"
        ws.append(["如何获取EventID"])
        ws.append(["EventID 通过 BMC 命令行查询获取。"])
        _save_wb(wb, self.imports / "说明.xlsx")
        body = self._doc_body("说明.xlsx")
        self.assertIn("如何获取EventID\nEventID 通过 BMC 命令行查询获取。", body)
        self.assertNotIn("[表头]", body)
        self.assertNotIn("=", body)

    def test_long_first_row_no_kv(self):
        """首行含长文本 cell（>30 字符）：是内容不是表头 → 裸拼。"""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["这是一段很长的说明文字" * 5, "b", "c"])
        ws.append(["d", "e", "f"])
        ws.append(["g", "h", "i"])
        _save_wb(wb, self.imports / "长文.xlsx")
        body = self._doc_body("长文.xlsx")
        self.assertNotIn("[表头]", body)
        self.assertIn("g h i", body)

    def test_ragged_rows_kv_with_tail(self):
        """数据行 cell 数超过表头：键值对截齐 + 多余 cell 裸值追加。"""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["a", "b"])
        ws.append(["1", "2", "3"])
        ws.append(["4", "5"])
        _save_wb(wb, self.imports / "ragged.xlsx")
        body = self._doc_body("ragged.xlsx")
        self.assertIn("a=1 | b=2 | 3", body)
        self.assertIn("a=4 | b=5", body)


if __name__ == "__main__":
    unittest.main()
